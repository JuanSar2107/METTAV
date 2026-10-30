"""Tests de la descarga en Excel y PDF del historial y de usuarios."""

from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from mi_app.extensions import db
from mi_app.models import MovimientoInventario


@pytest.fixture()
def admin(creador_usuarios, iniciar_sesion):
    creador_usuarios("admin", es_admin=True, nombre_completo="Jefa de Bodega")
    iniciar_sesion("admin")


@pytest.fixture()
def movimientos(app):
    with app.app_context():
        for fecha, tipo, articulo, codigo, categoria, formulario in [
            # 27/10 a las 23:30 en Colombia es el 28/10 a las 04:30 en UTC.
            (datetime(2026, 10, 28, 4, 30), "Retiro", "Banco hidráulico", "ZF-MAQ-CO-001", "Maquinaria", "1383926"),
            (datetime(2026, 10, 26, 15, 0), "Ingreso", "Sillas de oficina", "ZF-MO-CO-013", "Mobiliario", None),
        ]:
            db.session.add(MovimientoInventario(
                empleado="Jefa de Bodega", usuario="admin", articulo=articulo, codigo_articulo=codigo,
                categoria=categoria, tipo=tipo, cantidad=1, unidad="unidades", fecha_hora=fecha,
                formulario=formulario, origen="Manual",
            ))
        db.session.commit()


def _excel(respuesta):
    return list(load_workbook(BytesIO(respuesta.data)).active.iter_rows(values_only=True))


def _texto_pdf(respuesta):
    return "\n".join(pagina.extract_text() for pagina in PdfReader(BytesIO(respuesta.data)).pages)


def test_historial_en_excel_con_filtros(client, admin, movimientos):
    respuesta = client.get("/api/inventario/movimientos?formato=xlsx")
    assert respuesta.status_code == 200
    assert "historial-movimientos-" in respuesta.headers["Content-Disposition"]
    filas = _excel(respuesta)
    assert filas[0][:3] == ("Fecha (hora Colombia)", "Tipo", "Código")
    assert len(filas) == 3

    # Mismos filtros que la pantalla: búsqueda sin tildes, categoría y día en hora de Colombia.
    assert [f[2] for f in _excel(client.get("/api/inventario/movimientos?formato=xlsx&q=hidraulico"))[1:]] == ["ZF-MAQ-CO-001"]
    assert len(_excel(client.get("/api/inventario/movimientos?formato=xlsx&categoria=Mobiliario"))) == 2
    por_dia = _excel(client.get("/api/inventario/movimientos?formato=xlsx&dia=2026-10-27"))
    assert [(f[0], f[8]) for f in por_dia[1:]] == [("2026-10-27 23:30", "1383926")]


def test_historial_en_pdf(client, admin, movimientos):
    respuesta = client.get("/api/inventario/movimientos?formato=pdf&q=1383926")
    assert respuesta.status_code == 200
    assert respuesta.mimetype == "application/pdf"
    assert respuesta.data.startswith(b"%PDF")
    texto = _texto_pdf(respuesta)
    assert "Historial de movimientos" in texto
    assert "Banco hidráulico" in texto
    assert "Sillas de oficina" not in texto
    assert "Jefa de Bodega" in texto
    assert 'búsqueda "1383926"' in texto


def test_historial_sin_formato_sigue_devolviendo_json(client, admin, movimientos):
    datos = client.get("/api/inventario/movimientos").get_json()
    assert datos["total"] == 2


def test_usuarios_en_excel_y_pdf_con_filtros(client, admin, creador_usuarios):
    creador_usuarios("ana", nombre_completo="Ana Pérez", correo="ana@metttav.co", telefono="3001234567")
    creador_usuarios("luis", nombre_completo="Luis Gómez")

    filas = _excel(client.get("/api/usuarios?formato=xlsx"))
    assert filas[0] == ("Usuario", "Nombre completo", "Rol", "Correo", "Número de identidad", "Teléfono", "Dirección")
    assert sorted(f[0] for f in filas[1:]) == ["admin", "ana", "luis"]
    assert [f[0] for f in _excel(client.get("/api/usuarios?formato=xlsx&rol=admin"))[1:]] == ["admin"]
    assert [f[0] for f in _excel(client.get("/api/usuarios?formato=xlsx&q=3001234567"))[1:]] == ["ana"]

    texto = _texto_pdf(client.get("/api/usuarios?formato=pdf&rol=empleado"))
    assert "Usuarios y empleados" in texto
    assert "Ana Pérez" in texto and "Luis Gómez" in texto
    assert "Jefa de Bodega" in texto  # quién lo generó
    assert "solo empleados / usuarios" in texto


def test_descargas_requieren_admin(client, creador_usuarios, iniciar_sesion):
    assert client.get("/api/usuarios?formato=pdf").status_code == 401
    creador_usuarios("normal")
    iniciar_sesion("normal")
    assert client.get("/api/usuarios?formato=xlsx").status_code == 403
    assert client.get("/api/inventario/movimientos?formato=pdf").status_code == 403


def test_pdf_sin_registros(client, admin):
    respuesta = client.get("/api/inventario/movimientos?formato=pdf")
    assert respuesta.status_code == 200
    assert "No hay registros con esos filtros." in _texto_pdf(respuesta)


def test_dashboard_tiene_botones_de_descarga(client, admin):
    html = client.get("/dashboard").get_data(as_text=True)
    assert 'data-export-scope="movimientos"' in html
    assert 'data-export-scope="usuarios"' in html
    assert html.count('data-export-format="pdf"') == 2

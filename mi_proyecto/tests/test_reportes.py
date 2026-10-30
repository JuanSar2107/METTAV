"""Tests de reportes y conciliación con Sizfra (R7)."""

from datetime import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from mi_app.extensions import db
from mi_app.models import ArticuloInventario, MovimientoInventario


@pytest.fixture()
def admin(creador_usuarios, iniciar_sesion):
    creador_usuarios("admin", es_admin=True)
    iniciar_sesion("admin")


@pytest.fixture()
def inventario(app):
    datos = [
        ("ZF-MAQ-CO-001", "Banco hidráulico", "002", 2, Decimal("80.00")),
        ("ZF-MAQ-CO-002", "Banco de cubo", "002", 1, None),
        ("ZF-MO-CO-005", "Escritorio en L", "001", 2, Decimal("150.50")),
        ("HER-001", "Taladro", None, 12, None),
    ]
    with app.app_context():
        for codigo, nombre, grupo, stock, fob in datos:
            db.session.add(ArticuloInventario(
                codigo=codigo, nombre=nombre, categoria="Maquinaria", descripcion=nombre, grupo=grupo,
                stock=stock, stock_maximo=stock, unidad="unidades", ubicacion="Bodega",
                imagen_url="x.png", valor_fob_usd=fob,
            ))
        # 27/10 a las 23:30 en Colombia es el 28/10 a las 04:30 en UTC.
        for fecha, tipo, formulario in [
            (datetime(2026, 10, 28, 4, 30), "Retiro", "1383926"),
            (datetime(2026, 10, 26, 15, 0), "Ingreso", None),
        ]:
            db.session.add(MovimientoInventario(
                empleado="Jefe", usuario="admin", articulo="Banco hidráulico", codigo_articulo="ZF-MAQ-CO-001",
                categoria="Maquinaria", tipo=tipo, cantidad=1, unidad="unidades", fecha_hora=fecha,
                formulario=formulario,
            ))
        db.session.commit()


def test_reportes_requieren_admin(client, creador_usuarios, iniciar_sesion):
    assert client.get("/api/reportes/inventario").status_code == 401
    creador_usuarios("normal")
    iniciar_sesion("normal")
    for ruta in ("/api/reportes/inventario", "/api/reportes/movimientos", "/api/reportes/salidas-integradas"):
        assert client.get(ruta).status_code == 403
    assert client.post("/api/reportes/conciliacion").status_code == 403


def test_reporte_inventario_por_grupo_con_valor_total(client, admin, inventario):
    datos = client.get("/api/reportes/inventario?grupo=002").get_json()
    assert [f["codigo"] for f in datos["filas"]] == ["ZF-MAQ-CO-001", "ZF-MAQ-CO-002"]
    assert datos["filas"][0]["valor_total_usd"] == 160.0
    assert datos["resumen"] == [{
        "grupo": "002", "nombre": "Equipos de bodega", "articulos": 2, "unidades": 3, "valor_total_usd": 160.0,
    }]
    assert client.get("/api/reportes/inventario?grupo=sin-grupo").get_json()["total"] == 1
    assert client.get("/api/reportes/inventario?q=escritorio").get_json()["filas"][0]["codigo"] == "ZF-MO-CO-005"


def test_reporte_inventario_en_excel(client, admin, inventario):
    respuesta = client.get("/api/reportes/inventario?formato=xlsx")
    assert respuesta.status_code == 200
    assert "reporte-inventario-" in respuesta.headers["Content-Disposition"]
    hoja = load_workbook(BytesIO(respuesta.data)).active
    filas = list(hoja.iter_rows(values_only=True))
    assert filas[0][:3] == ("Código ZF", "Código interno", "Descripción")
    assert len(filas) == 5
    escritorio = next(f for f in filas if f[0] == "ZF-MO-CO-005")
    assert escritorio[11] == 301.0  # valor total USD = 2 x 150,50


def test_reporte_movimientos_usa_dias_de_colombia(client, admin, inventario):
    datos = client.get("/api/reportes/movimientos?desde=2026-10-27&hasta=2026-10-27").get_json()
    assert [(f["tipo"], f["formulario"]) for f in datos["filas"]] == [("Retiro", "1383926")]
    assert datos["resumen"]["unidades_retiradas"] == 1
    todos = client.get("/api/reportes/movimientos").get_json()
    assert todos["total"] == 2
    assert client.get("/api/reportes/movimientos?tipo=Ingreso").get_json()["total"] == 1


def test_reporte_movimientos_valida_fechas(client, admin):
    assert client.get("/api/reportes/movimientos?desde=27-10-2026").status_code == 400
    assert client.get("/api/reportes/movimientos?desde=2026-10-28&hasta=2026-10-27").status_code == 400


def test_reporte_salidas_integradas(client, app, admin, inventario):
    with app.app_context():
        banco = ArticuloInventario.query.filter_by(codigo="ZF-MAQ-CO-001").one().id
        parte = ArticuloInventario.query.filter_by(codigo="ZF-MAQ-CO-002").one().id
    client.post("/api/inventario/salidas-integradas", json={
        "principal": {"id": banco, "cantidad": 1},
        "componentes": [{"id": parte, "cantidad": 1}],
        "proceso": "Fabricación",
    })
    datos = client.get("/api/reportes/salidas-integradas").get_json()
    assert datos["total"] == 1
    assert [m["rol"] for m in datos["salidas"][0]["movimientos"]] == ["Principal", "Integrado"]
    hoja = load_workbook(BytesIO(client.get("/api/reportes/salidas-integradas?formato=xlsx").data)).active
    assert hoja.max_row == 3  # encabezado + equipo + parte


def _archivo_sizfra(filas, encabezados=("FORMULARIO", "CODIGO PRODUCTO", "DESCRIPCION", "SALDO")):
    libro = Workbook()
    hoja = libro.active
    hoja.append(["Reporte 026 - Saldos consolidados"])  # título antes de los encabezados
    hoja.append(list(encabezados))
    for fila in filas:
        hoja.append(fila)
    bufer = BytesIO()
    libro.save(bufer)
    bufer.seek(0)
    return bufer


def _conciliar(client, archivo, nombre="saldos.xlsx", **datos):
    return client.post(
        "/api/reportes/conciliacion",
        data={"archivo": (archivo, nombre), **datos},
        content_type="multipart/form-data",
    )


def test_conciliacion_marca_diferencias(client, admin, inventario):
    archivo = _archivo_sizfra([
        ["1383926", "ZF-MAQ-CO-001", "Banco", "2,000000"],   # coincide (formato de Sizfra)
        ["1383927", "ZF-MAQ-CO-002", "Banco cubo", "1"],
        ["1383928", "ZF-MAQ-CO-002", "Banco cubo", "1"],     # se suma: 2 en Sizfra, 1 en la app
        ["1383929", "ZF-OT-CO-099", "Algo", "4"],            # no existe en la app
        [None, "TOTAL", None, "8"],                          # fila de totales, se ignora
    ])
    datos = _conciliar(client, archivo).get_json()
    assert datos["columnas"] == {"codigo": "CODIGO PRODUCTO", "saldo": "SALDO"}
    estados = {f["codigo"]: (f["estado"], f["diferencia"]) for f in datos["filas"]}
    assert estados["ZF-MAQ-CO-001"] == ("Coincide", 0.0)
    assert estados["ZF-MAQ-CO-002"] == ("Diferencia", -1.0)
    assert estados["ZF-OT-CO-099"] == ("Solo en Sizfra", None)
    assert estados["ZF-MO-CO-005"] == ("Solo en la app", None)
    # El taladro de demostración no tiene código ZF: no se compara por defecto.
    assert "HER-001" not in estados
    assert datos["filas"][0]["estado"] == "Diferencia"  # lo que hay que revisar va primero
    assert datos["resumen"] == {"Diferencia": 1, "Solo en Sizfra": 1, "Solo en la app": 1, "Coincide": 1}


def test_conciliacion_con_csv_y_todos_los_codigos(client, admin, inventario):
    csv = "Codigo;Cantidad\nZF-MAQ-CO-001;2\nHER-001;12\n".encode("latin-1")
    datos = _conciliar(client, BytesIO(csv), "saldos.csv", solo_zf="0").get_json()
    estados = {f["codigo"]: f["estado"] for f in datos["filas"]}
    assert estados["HER-001"] == "Coincide"
    assert estados["ZF-MAQ-CO-002"] == "Solo en la app"


def test_conciliacion_en_excel(client, admin, inventario):
    respuesta = _conciliar(client, _archivo_sizfra([["1", "ZF-MAQ-CO-001", "Banco", "3"]]), formato="xlsx")
    hoja = load_workbook(BytesIO(respuesta.data)).active
    assert hoja.cell(row=2, column=6).value == "Diferencia"


def test_conciliacion_archivo_sin_columnas(client, admin):
    respuesta = _conciliar(client, _archivo_sizfra([["a", "b"]], encabezados=("Nombre", "Precio")))
    assert respuesta.status_code == 400
    assert "columnas de código y saldo" in respuesta.get_json()["error"]
    assert _conciliar(client, BytesIO(b"x"), "saldos.pdf").status_code == 400


def test_dashboard_tiene_pestana_de_reportes(client, admin):
    html = client.get("/dashboard").get_data(as_text=True)
    assert 'data-target="section-reports"' in html
    assert 'data-report="conciliacion"' in html

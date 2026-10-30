"""Campos de zona franca de los artículos e importación desde Excel.

El Excel `Inventario_Consecutivo_Por_Categoria EQUIPOS METTAV` trae por cada
ítem el código interno, el código ZF (el que se registra en Sizfra), la
categoría, la ubicación, el tipo, el consecutivo, la descripción y la marca.
"""

import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from mi_app.models.inventory import GRUPOS_INVENTARIO, PROCEDENCIAS

# Siglas que usa el Excel en las columnas de categoría y ubicación.
CATEGORIAS_SIGLA = {
    "EQ": "Equipos",
    "MAQ": "Maquinaria",
    "EL": "Electrodomésticos",
    "MO": "Mobiliario",
    "HE": "Herramientas",
    "OT": "Otros",
}
UBICACIONES_SIGLA = {
    "ADM": "Administración",
    "BOD": "Bodega",
    "COM": "Comedor",
    "OP": "Operativa",
    "APT": "Apartamento",
}

_RE_CODIGO = re.compile(r"^[A-Z0-9][A-Z0-9-]*$")
_RE_SUBPARTIDA = re.compile(r"^\d{4}(\.?\d{2}){0,3}$")
_RE_MILES = re.compile(r"^\d{1,3}(\.\d{3})+$")
_RE_UNIDADES = re.compile(r"\((\d+)\s*(?:und|unds|unidades)\.?\)", re.IGNORECASE)


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def leer_valor_fob(texto: str) -> Decimal | None:
    """Convierte '1200', '1.200,50' o '1200.50' a Decimal; vacío devuelve None."""
    texto = texto.replace("$", "").replace("USD", "").replace(" ", "").strip()
    if not texto:
        return None
    if "," in texto and "." in texto:
        # El último separador es el decimal.
        if texto.rfind(",") > texto.rfind("."):
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
    elif "," in texto:
        texto = texto.replace(",", ".")
    elif _RE_MILES.match(texto):
        texto = texto.replace(".", "")
    try:
        valor = Decimal(texto)
    except InvalidOperation:
        raise ValueError("El valor FOB debe ser un número en dólares, por ejemplo 1200 o 1200,50.")
    if valor < 0 or valor >= Decimal("1e12"):
        raise ValueError("El valor FOB debe ser un número positivo.")
    return valor.quantize(Decimal("0.01"))


def campos_zona_franca(datos) -> dict:
    """Valida los campos opcionales de zona franca de un formulario o fila.

    Devuelve un dict listo para asignar al artículo. Lanza ValueError con un
    mensaje para el usuario si algún valor no es válido.
    """
    codigo_interno = _texto(datos.get("codigo_interno")).upper()
    grupo = _texto(datos.get("grupo"))
    tipo = _texto(datos.get("tipo")).upper()
    marca = _texto(datos.get("marca"))
    subpartida = _texto(datos.get("subpartida"))
    procedencia = _texto(datos.get("procedencia")).capitalize()
    estado = _texto(datos.get("estado"))

    if codigo_interno and (len(codigo_interno) > 32 or not _RE_CODIGO.match(codigo_interno)):
        raise ValueError("El código interno solo admite letras, números y guiones (máximo 32).")
    if grupo and grupo.zfill(3) not in GRUPOS_INVENTARIO:
        raise ValueError("El grupo de inventario debe ser 001, 002, 003 o 004.")
    if len(tipo) > 8:
        raise ValueError("El tipo no puede superar 8 caracteres.")
    if len(marca) > 80:
        raise ValueError("La marca no puede superar 80 caracteres.")
    if marca.upper() == "N/A":
        marca = ""
    if subpartida and not _RE_SUBPARTIDA.match(subpartida):
        raise ValueError("La subpartida arancelaria debe tener el formato 8413.30.10.00.")
    if procedencia and procedencia not in PROCEDENCIAS:
        raise ValueError("La procedencia debe ser Nacional o Extranjera.")
    if len(estado) > 40:
        raise ValueError("El estado no puede superar 40 caracteres.")

    return {
        "codigo_interno": codigo_interno or None,
        "grupo": grupo.zfill(3) if grupo else None,
        "tipo": tipo or None,
        "marca": marca or None,
        "subpartida": subpartida or None,
        "valor_fob_usd": leer_valor_fob(_texto(datos.get("valor_fob_usd"))),
        "procedencia": procedencia or None,
        "estado": estado or None,
    }


_MESES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7,
    "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12,
}
_RE_FECHA_TEXTO = re.compile(r"^(\d{1,2})\s+de\s+([a-z]+)\s+(?:de|del)\s+(\d{4})$")


def leer_fecha(valor) -> date | None:
    """Acepta fechas de Excel, '2026-05-08', '08/05/2026' u '08 de mayo del 2026'."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = unicodedata.normalize("NFKD", _texto(valor)).encode("ascii", "ignore").decode().lower()
    texto = " ".join(texto.split())
    if not texto:
        return None
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            pass
    coincidencia = _RE_FECHA_TEXTO.match(texto)
    if coincidencia and coincidencia.group(2) in _MESES:
        dia, mes, anio = int(coincidencia.group(1)), _MESES[coincidencia.group(2)], int(coincidencia.group(3))
        try:
            return date(anio, mes, dia)
        except ValueError:
            pass
    raise ValueError(f"La fecha '{_texto(valor)}' no es válida. Usa 08/05/2026 o '8 de mayo del 2026'.")


def campos_trazabilidad(datos) -> dict:
    """Valida serial, responsable, fechas y documentos soporte (todos opcionales)."""
    serial = _texto(datos.get("serial"))
    if serial.upper() == "N/A":
        serial = ""
    textos = {
        "serial": (serial, 80),
        "responsable": (_texto(datos.get("responsable")), 120),
        "documento_soporte": (_texto(datos.get("documento_soporte")), 160),
        "documento_retiro": (_texto(datos.get("documento_retiro")), 160),
        "razon_retiro": (_texto(datos.get("razon_retiro")), 2000),
    }
    campos = {}
    for campo, (valor, maximo) in textos.items():
        if len(valor) > maximo:
            raise ValueError(f"El campo {campo.replace('_', ' ')} no puede superar {maximo} caracteres.")
        campos[campo] = valor or None
    campos["fecha_ingreso"] = leer_fecha(datos.get("fecha_ingreso"))
    campos["fecha_retiro"] = leer_fecha(datos.get("fecha_retiro"))
    return campos


def normalizar_codigo(texto: str) -> str:
    return _texto(texto).upper()


def codigo_valido(codigo: str) -> bool:
    return bool(codigo) and len(codigo) <= 24 and bool(_RE_CODIGO.match(codigo))


# --- Importación desde Excel -------------------------------------------------

def _clave(encabezado) -> str:
    """'Código ZF' -> 'codigo zf'; 'Categoria (EQ/MAQ/...)' -> 'categoria'."""
    texto = unicodedata.normalize("NFKD", _texto(encabezado)).encode("ascii", "ignore").decode()
    texto = re.sub(r"\(.*?\)", " ", texto.lower())
    return " ".join(re.sub(r"[^a-z0-9]+", " ", texto).split())


# Columna de la app -> nombres posibles del encabezado en el Excel.
_ENCABEZADOS = {
    "codigo_interno": ("codigo interno",),
    "codigo": ("codigo zf", "codigo", "codigo producto"),
    "categoria": ("categoria",),
    "ubicacion": ("ubicacion",),
    "tipo": ("tipo",),
    "consecutivo": ("consecutivo",),
    "descripcion": ("descripcion", "descripcion del producto"),
    "marca": ("marca",),
    "grupo": ("grupo", "grupo de inventario"),
    "subpartida": ("subpartida", "subpartida arancelaria", "sub partida"),
    # En el Excel de METTAV "Valor FOB" está en pesos y "valor USD" en dólares:
    # solo se usa "Valor FOB" si no hay una columna explícita en dólares.
    "valor_fob_usd": ("valor fob usd", "valor usd", "valor fob"),
    "procedencia": ("procedencia",),
    "estado": ("estado",),
    "stock": ("cantidad", "cant", "stock"),
    "unidad": ("unidad", "unidad de medida"),
    "serial": ("serial", "numero de serie"),
    "responsable": ("responsable",),
    "fecha_ingreso": ("fecha de ingreso", "fecha ingreso"),
    "documento_soporte": ("documento soporte", "documento de ingreso"),
    "fecha_retiro": ("fecha de retiro", "fecha retiro"),
    # El Excel tiene dos columnas "Documento soporte": la primera es la del
    # ingreso y la segunda (después de "Fecha de Retiro") la del retiro.
    "documento_retiro": ("documento de retiro", "documento soporte retiro", "documento soporte"),
    "razon_retiro": ("razones por las cuales se retira", "razon de retiro", "razon del retiro"),
}


def _mapa_columnas(fila) -> dict:
    claves = [_clave(celda) for celda in fila]
    mapa = {}
    for campo, opciones in _ENCABEZADOS.items():
        # Las opciones van en orden de preferencia.
        for opcion in opciones:
            indice = next(
                (i for i, clave in enumerate(claves) if clave == opcion and i not in mapa.values()),
                None,
            )
            if indice is not None:
                mapa[campo] = indice
                break
    return mapa


def grupo_por_ubicacion(ubicacion_sigla: str) -> str:
    """Bodega -> 002, operativa -> 003, el resto (oficina, comedor, apto.) -> 001."""
    if ubicacion_sigla == "BOD":
        return "002"
    if ubicacion_sigla in ("OP", "M/OP"):
        return "003"
    return "001"


def leer_excel(archivo) -> tuple[list[dict], list[dict], list[dict]]:
    """Lee la hoja de inventario y devuelve (filas válidas, errores, avisos).

    Un error deja la fila fuera; un aviso (p. ej. una fecha ilegible) la
    importa igual pero sin ese dato.

    Las columnas de código del Excel son fórmulas: se usa el último valor que
    guardó Excel y, si no lo hay, se arma el código ZF con la misma fórmula
    ("ZF-" & categoría & "-" & tipo & "-" & consecutivo en 3 dígitos).
    """
    from openpyxl import load_workbook

    try:
        libro = load_workbook(archivo, read_only=True, data_only=True)
    except Exception:
        raise ValueError("El archivo no es un Excel (.xlsx) válido.")

    hoja = libro["Inventario Equipos"] if "Inventario Equipos" in libro.sheetnames else libro.worksheets[0]
    filas = hoja.iter_rows(values_only=True)

    mapa = {}
    numero_encabezado = 0
    for numero, fila in enumerate(filas, start=1):
        mapa = _mapa_columnas(fila)
        if "descripcion" in mapa and ("codigo" in mapa or "consecutivo" in mapa):
            numero_encabezado = numero
            break
        if numero >= 15:
            break
    if not numero_encabezado:
        libro.close()
        raise ValueError(
            "No se encontró la fila de encabezados. El Excel debe tener al menos las columnas "
            "'Código ZF' (o 'Consecutivo') y 'Descripción'."
        )

    validas, errores, avisos = [], [], []
    for numero, fila in enumerate(filas, start=numero_encabezado + 1):
        def crudo(campo):
            indice = mapa.get(campo)
            return fila[indice] if indice is not None and indice < len(fila) else None

        def celda(campo):
            return _texto(crudo(campo))

        descripcion = celda("descripcion")
        if not descripcion:
            continue  # filas vacías o separadores

        categoria_sigla = celda("categoria").upper()
        ubicacion_sigla = celda("ubicacion").upper()
        tipo = celda("tipo").upper()
        codigo = normalizar_codigo(celda("codigo"))
        if not codigo and categoria_sigla and tipo and celda("consecutivo").isdigit():
            codigo = f"ZF-{categoria_sigla}-{tipo}-{int(celda('consecutivo')):03d}"
        if not codigo_valido(codigo):
            errores.append({"fila": numero, "error": "Falta el código ZF o no es válido."})
            continue

        try:
            campos = campos_zona_franca({
                "codigo_interno": celda("codigo_interno"),
                "grupo": celda("grupo") or grupo_por_ubicacion(ubicacion_sigla),
                "tipo": tipo,
                "marca": celda("marca"),
                "subpartida": celda("subpartida"),
                "valor_fob_usd": celda("valor_fob_usd"),
                "procedencia": celda("procedencia"),
                "estado": celda("estado"),
            })
        except ValueError as error:
            errores.append({"fila": numero, "codigo": codigo, "error": str(error)})
            continue

        datos_trazabilidad = {
            campo: crudo(campo)
            for campo in ("serial", "responsable", "fecha_ingreso", "documento_soporte",
                          "fecha_retiro", "documento_retiro", "razon_retiro")
        }
        for campo_fecha in ("fecha_ingreso", "fecha_retiro"):
            try:
                leer_fecha(datos_trazabilidad[campo_fecha])
            except ValueError as error:
                avisos.append({"fila": numero, "codigo": codigo, "aviso": str(error)})
                datos_trazabilidad[campo_fecha] = None
        try:
            campos.update(campos_trazabilidad(datos_trazabilidad))
        except ValueError as error:
            errores.append({"fila": numero, "codigo": codigo, "error": str(error)})
            continue

        stock_texto = celda("stock")
        if stock_texto.isdigit():
            stock = int(stock_texto)
        else:
            # "sillones color gris (2 und)" -> 2; si no dice nada, 1 unidad.
            coincidencia = _RE_UNIDADES.search(descripcion)
            stock = int(coincidencia.group(1)) if coincidencia else 1

        validas.append({
            "fila": numero,
            "codigo": codigo,
            "nombre": descripcion[:120],
            "descripcion": descripcion[:4000],
            "categoria": CATEGORIAS_SIGLA.get(categoria_sigla, categoria_sigla.title() or "General")[:60],
            "ubicacion": UBICACIONES_SIGLA.get(ubicacion_sigla, ubicacion_sigla or "Sin asignar")[:80],
            "unidad": (celda("unidad") or "unidades")[:24].lower(),
            "stock": stock,
            **campos,
        })

    libro.close()
    return validas, errores, avisos

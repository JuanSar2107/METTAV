"""Descarga de listados en Excel y PDF (reportes, historial y usuarios)."""

from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from flask import current_app, send_file

_HORA_COLOMBIA = timezone(timedelta(hours=-5))
_AZUL = "054A86"


def _valor(valor):
    """Decimal -> float para Excel (los valores FOB tienen 2 decimales)."""
    return float(valor) if isinstance(valor, Decimal) else valor


def _nombre(base: str, extension: str) -> str:
    return f"{base}-{date.today().isoformat()}.{extension}"


def respuesta_excel(titulo: str, encabezados: list[str], filas: list[list], nombre_archivo: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=_AZUL)
    for fila in filas:
        hoja.append([_valor(valor) for valor in fila])
    for indice, encabezado in enumerate(encabezados, start=1):
        largo = max([len(str(encabezado))] + [len(str(f[indice - 1] or "")) for f in filas[:500]])
        hoja.column_dimensions[hoja.cell(row=1, column=indice).column_letter].width = min(max(10, largo + 2), 60)
    hoja.freeze_panes = "A2"

    bufer = BytesIO()
    libro.save(bufer)
    bufer.seek(0)
    return send_file(
        bufer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=_nombre(nombre_archivo, "xlsx"),
    )


def respuesta_pdf(
    titulo: str,
    encabezados: list[str],
    filas: list[list],
    nombre_archivo: str,
    anchos: list[float],
    detalle: str = "",
    generado_por: str = "",
):
    """Tabla en PDF horizontal (carta) con el encabezado de METTAV.

    `anchos` son proporciones de cada columna (suman 1). `detalle` describe los
    filtros aplicados.
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    azul = colors.HexColor("#" + _AZUL)
    gris = colors.HexColor("#526B80")
    base = getSampleStyleSheet()
    estilo_celda = ParagraphStyle("celda", parent=base["Normal"], fontSize=8, leading=10)
    estilo_titulo_col = ParagraphStyle("col", parent=estilo_celda, fontName="Helvetica-Bold", textColor=colors.white)
    estilo_titulo = ParagraphStyle("titulo", parent=base["Title"], fontSize=17, leading=21, textColor=azul,
                                   alignment=0, spaceAfter=2)
    estilo_info = ParagraphStyle("info", parent=base["Normal"], fontSize=8.5, leading=11, textColor=gris)

    tamano = landscape(letter)
    margen = 0.5 * inch
    ancho_util = tamano[0] - 2 * margen
    ahora = datetime.now(_HORA_COLOMBIA).strftime("%d/%m/%Y %H:%M")

    def texto(valor) -> str:
        if valor is None:
            return ""
        if isinstance(valor, Decimal):
            valor = f"{valor:,.2f}"
        return (str(valor).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

    historia = []
    logo = Path(current_app.static_folder or "") / "images" / "mettav-logo.png"
    cabecera = [Paragraph(f"METTAV GROUP S.A.S.<br/>{texto(titulo)}", estilo_titulo)]
    if logo.is_file():
        imagen = Image(str(logo), width=0.6 * inch, height=0.6 * inch, kind="proportional")
        cabecera = Table([[imagen, cabecera[0]]], colWidths=[0.75 * inch, ancho_util - 0.75 * inch])
        cabecera.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
        historia.append(cabecera)
    else:
        historia += cabecera
    info = f"Generado el {ahora} (hora de Colombia)"
    if generado_por:
        info += f" por {texto(generado_por)}"
    info += f" · {len(filas)} registros"
    if detalle:
        info += f"<br/>Filtros: {texto(detalle)}"
    historia += [Paragraph(info, estilo_info), Spacer(1, 8)]

    datos = [[Paragraph(texto(c), estilo_titulo_col) for c in encabezados]]
    datos += [[Paragraph(texto(v), estilo_celda) for v in fila] for fila in filas]
    if not filas:
        datos.append([Paragraph("No hay registros con esos filtros.", estilo_celda)] + [""] * (len(encabezados) - 1))
    tabla = Table(datos, colWidths=[ancho_util * a for a in anchos], repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), azul),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C7D3")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EDF4F8")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    historia.append(tabla)

    def pie(canvas, documento):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(gris)
        canvas.drawString(margen, 0.3 * inch, f"METTAV GROUP S.A.S. · {titulo}")
        canvas.drawRightString(tamano[0] - margen, 0.3 * inch, f"Página {documento.page}")
        canvas.restoreState()

    bufer = BytesIO()
    documento = SimpleDocTemplate(
        bufer, pagesize=tamano, leftMargin=margen, rightMargin=margen, topMargin=margen, bottomMargin=0.55 * inch,
        title=f"{titulo} - METTAV", author="METTAV Group S.A.S.",
    )
    documento.build(historia, onFirstPage=pie, onLaterPages=pie)
    bufer.seek(0)
    return send_file(bufer, mimetype="application/pdf", as_attachment=True,
                     download_name=_nombre(nombre_archivo, "pdf"))

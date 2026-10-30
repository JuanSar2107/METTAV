"""Reportes del inventario y conciliación con Sizfra (R7).

Cada reporte responde JSON y, con ?formato=xlsx, el mismo contenido como
archivo de Excel para las auditorías de la zona franca.
"""

from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from io import BytesIO

from flask import Blueprint, jsonify, request, send_file

from mi_app.models import ArticuloInventario, MovimientoInventario, SalidaIntegrada
from mi_app.models.inventory import GRUPOS_INVENTARIO
from mi_app.security import usuario_actual
from mi_app.zona_franca import leer_saldos_sizfra

reportes_bp = Blueprint("reportes", __name__, url_prefix="/api/reportes")

# Colombia no tiene horario de verano: UTC-5 todo el año. Las fechas se
# guardan en UTC y los filtros "desde/hasta" son días de Colombia.
_HORA_COLOMBIA = timezone(timedelta(hours=-5))


def _error_admin():
    usuario = usuario_actual()
    if usuario is None:
        return jsonify(error="Debes iniciar sesión."), 401
    if not usuario.es_admin:
        return jsonify(error="Solo un administrador puede ver los reportes."), 403
    return None


def _fecha_param(nombre: str) -> date | None:
    texto = (request.args.get(nombre) or "").strip()
    if not texto:
        return None
    try:
        return date.fromisoformat(texto)
    except ValueError:
        raise ValueError(f"La fecha '{nombre}' debe tener el formato AAAA-MM-DD.")


def _rango_utc() -> tuple[datetime | None, datetime | None]:
    """Convierte ?desde y ?hasta (días en Colombia, ambos incluidos) a UTC sin zona."""
    desde, hasta = _fecha_param("desde"), _fecha_param("hasta")
    if desde and hasta and desde > hasta:
        raise ValueError("La fecha 'desde' no puede ser posterior a 'hasta'.")

    def a_utc(dia: date) -> datetime:
        return datetime.combine(dia, time.min, _HORA_COLOMBIA).astimezone(timezone.utc).replace(tzinfo=None)

    return (a_utc(desde) if desde else None, a_utc(hasta + timedelta(days=1)) if hasta else None)


def _hora_local(valor: datetime) -> str:
    return valor.replace(tzinfo=timezone.utc).astimezone(_HORA_COLOMBIA).strftime("%Y-%m-%d %H:%M")


def _numero(valor):
    """Decimal -> float para JSON y Excel (los valores FOB tienen 2 decimales)."""
    return float(valor) if isinstance(valor, Decimal) else valor


def _excel(titulo: str, encabezados: list[str], filas: list[list], nombre_archivo: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="054A86")
    for fila in filas:
        hoja.append([_numero(valor) for valor in fila])
    for indice, encabezado in enumerate(encabezados, start=1):
        largo = max([len(str(encabezado))] + [len(str(f[indice - 1] or "")) for f in filas[:500]])
        hoja.column_dimensions[hoja.cell(row=1, column=indice).column_letter].width = min(max(10, largo + 2), 60)
    hoja.freeze_panes = "A2"

    bufer = BytesIO()
    libro.save(bufer)
    bufer.seek(0)
    return send_file(
        bufer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"{nombre_archivo}-{date.today().isoformat()}.xlsx",
    )


def _quiere_excel() -> bool:
    return (request.args.get("formato") or request.form.get("formato") or "").lower() == "xlsx"


@reportes_bp.get("/inventario")
def reporte_inventario():
    """Inventario actual, filtrable por grupo de Sizfra, categoría o texto."""
    error = _error_admin()
    if error is not None:
        return error

    consulta = ArticuloInventario.query
    grupo = (request.args.get("grupo") or "").strip()
    categoria = (request.args.get("categoria") or "").strip()
    texto = (request.args.get("q") or "").strip()
    if grupo == "sin-grupo":
        consulta = consulta.filter(ArticuloInventario.grupo.is_(None))
    elif grupo:
        consulta = consulta.filter(ArticuloInventario.grupo == grupo)
    if categoria:
        consulta = consulta.filter(ArticuloInventario.categoria == categoria)
    estado_servicio = (request.args.get("estado_servicio") or "").strip()
    if estado_servicio:
        consulta = consulta.filter(ArticuloInventario.estado_servicio == estado_servicio)
    if texto:
        patron = f"%{texto}%"
        consulta = consulta.filter(
            ArticuloInventario.codigo.ilike(patron)
            | ArticuloInventario.codigo_interno.ilike(patron)
            | ArticuloInventario.nombre.ilike(patron)
        )
    articulos = consulta.order_by(ArticuloInventario.grupo, ArticuloInventario.codigo).all()

    filas = []
    resumen = {}
    for a in articulos:
        valor_total = a.valor_fob_usd * a.stock if a.valor_fob_usd is not None else None
        filas.append({
            "codigo": a.codigo,
            "codigo_interno": a.codigo_interno,
            "nombre": a.nombre,
            "grupo": a.grupo,
            "grupo_nombre": GRUPOS_INVENTARIO.get(a.grupo),
            "categoria": a.categoria,
            "ubicacion": a.ubicacion,
            "stock": a.stock,
            "unidad": a.unidad,
            "subpartida": a.subpartida,
            "valor_fob_usd": _numero(a.valor_fob_usd),
            "valor_total_usd": _numero(valor_total),
            "procedencia": a.procedencia,
            "estado": a.estado,
            "responsable": a.responsable,
            "serial": a.serial,
            "cliente": a.cliente,
            "tipo_servicio": a.tipo_servicio,
            "estado_servicio": a.estado_servicio,
        })
        clave = a.grupo or "Sin grupo"
        grupo_resumen = resumen.setdefault(clave, {
            "grupo": clave,
            "nombre": GRUPOS_INVENTARIO.get(a.grupo, "Sin grupo de Sizfra"),
            "articulos": 0,
            "unidades": 0,
            "valor_total_usd": Decimal("0"),
        })
        grupo_resumen["articulos"] += 1
        grupo_resumen["unidades"] += a.stock
        grupo_resumen["valor_total_usd"] += valor_total or 0
    for grupo_resumen in resumen.values():
        grupo_resumen["valor_total_usd"] = _numero(grupo_resumen["valor_total_usd"])

    if _quiere_excel():
        return _excel(
            "Inventario",
            ["Código ZF", "Código interno", "Descripción", "Grupo", "Nombre del grupo", "Categoría",
             "Ubicación", "Stock", "Unidad", "Subpartida", "Valor FOB USD", "Valor total USD",
             "Procedencia", "Estado", "Responsable", "Serial", "Cliente", "Servicio", "Estado del servicio"],
            [list(f.values()) for f in filas],
            "reporte-inventario",
        )
    return jsonify(total=len(filas), resumen=sorted(resumen.values(), key=lambda r: r["grupo"]), filas=filas)


@reportes_bp.get("/movimientos")
def reporte_movimientos():
    """Ingresos y retiros en un rango de fechas, con los datos de Sizfra."""
    error = _error_admin()
    if error is not None:
        return error
    try:
        desde, hasta = _rango_utc()
    except ValueError as error:
        return jsonify(error=str(error)), 400

    consulta = MovimientoInventario.query
    if desde:
        consulta = consulta.filter(MovimientoInventario.fecha_hora >= desde)
    if hasta:
        consulta = consulta.filter(MovimientoInventario.fecha_hora < hasta)
    tipo = (request.args.get("tipo") or "").strip()
    if tipo in ("Ingreso", "Retiro"):
        consulta = consulta.filter(MovimientoInventario.tipo == tipo)
    codigo = (request.args.get("codigo") or "").strip().upper()
    if codigo:
        consulta = consulta.filter(MovimientoInventario.codigo_articulo == codigo)
    movimientos = consulta.order_by(MovimientoInventario.fecha_hora, MovimientoInventario.id).all()

    filas = [m.to_dict() for m in movimientos]
    resumen = {
        "ingresos": sum(1 for m in movimientos if m.tipo == "Ingreso"),
        "retiros": sum(1 for m in movimientos if m.tipo == "Retiro"),
        "unidades_ingresadas": sum(m.cantidad for m in movimientos if m.tipo == "Ingreso"),
        "unidades_retiradas": sum(m.cantidad for m in movimientos if m.tipo == "Retiro"),
    }

    if _quiere_excel():
        return _excel(
            "Movimientos",
            ["Fecha (hora Colombia)", "Tipo", "Código", "Artículo", "Categoría", "Cantidad", "Unidad",
             "Operación Sizfra", "Formulario", "Origen", "Salida integrada", "Rol", "Referencia",
             "Observaciones", "Empleado", "Usuario"],
            [[_hora_local(m.fecha_hora), m.tipo, m.codigo_articulo, m.articulo, m.categoria, m.cantidad,
              m.unidad, m.operacion_sizfra, m.formulario, m.origen, m.id_salida, m.rol,
              m.salida.referencia if m.salida else None, m.observaciones, m.empleado, m.usuario]
             for m in movimientos],
            "reporte-movimientos",
        )
    return jsonify(total=len(filas), resumen=resumen, filas=filas)


@reportes_bp.get("/salidas-integradas")
def reporte_salidas_integradas():
    """Salidas integradas: qué partes se integraron a qué equipo y cuándo."""
    error = _error_admin()
    if error is not None:
        return error
    try:
        desde, hasta = _rango_utc()
    except ValueError as error:
        return jsonify(error=str(error)), 400

    consulta = SalidaIntegrada.query
    if desde:
        consulta = consulta.filter(SalidaIntegrada.fecha_hora >= desde)
    if hasta:
        consulta = consulta.filter(SalidaIntegrada.fecha_hora < hasta)
    salidas = consulta.order_by(SalidaIntegrada.fecha_hora, SalidaIntegrada.id).all()

    if _quiere_excel():
        return _excel(
            "Salidas integradas",
            ["Salida", "Fecha (hora Colombia)", "Proceso", "Referencia del equipo", "Rol", "Código",
             "Artículo", "Cantidad", "Unidad", "Operación Sizfra", "Formulario", "Observaciones", "Empleado"],
            [[s.id, _hora_local(s.fecha_hora), s.proceso, s.referencia, m.rol, m.codigo_articulo,
              m.articulo, m.cantidad, m.unidad, s.operacion_sizfra, s.formulario, s.observaciones, s.empleado]
             for s in salidas for m in s.movimientos],
            "reporte-salidas-integradas",
        )
    return jsonify(total=len(salidas), salidas=[s.to_dict() for s in salidas])


@reportes_bp.post("/conciliacion")
def conciliacion():
    """Compara el saldo de cada código en un reporte exportado de Sizfra con el stock de la app.

    Por defecto solo se comparan los artículos de la app con código ZF
    (los que existen en Sizfra); con solo_zf=0 se comparan todos.
    """
    error = _error_admin()
    if error is not None:
        return error

    archivo = request.files.get("archivo")
    if archivo is None or not archivo.filename:
        return jsonify(error="Selecciona el reporte de inventario exportado de Sizfra (.xlsx o .csv)."), 400
    try:
        saldos, columnas = leer_saldos_sizfra(archivo.stream, archivo.filename)
    except ValueError as error:
        return jsonify(error=str(error)), 400

    solo_zf = (request.form.get("solo_zf") or "1") != "0"
    articulos = {
        a.codigo: a for a in ArticuloInventario.query.all()
        if not solo_zf or a.codigo.startswith("ZF-")
    }
    if solo_zf:
        saldos = {codigo: saldo for codigo, saldo in saldos.items() if codigo.startswith("ZF-")}

    filas = []
    for codigo in sorted(set(saldos) | set(articulos)):
        articulo = articulos.get(codigo)
        saldo = saldos.get(codigo)
        stock = articulo.stock if articulo else None
        if articulo is None:
            estado = "Solo en Sizfra"
        elif saldo is None:
            estado = "Solo en la app"
        elif Decimal(stock) == saldo:
            estado = "Coincide"
        else:
            estado = "Diferencia"
        filas.append({
            "codigo": codigo,
            "nombre": articulo.nombre if articulo else None,
            "stock_app": stock,
            "saldo_sizfra": _numero(saldo),
            "diferencia": _numero(Decimal(stock) - saldo) if articulo and saldo is not None else None,
            "estado": estado,
        })
    # Primero lo que hay que revisar.
    orden = {"Diferencia": 0, "Solo en Sizfra": 1, "Solo en la app": 2, "Coincide": 3}
    filas.sort(key=lambda f: (orden[f["estado"]], f["codigo"]))
    resumen = {estado: sum(1 for f in filas if f["estado"] == estado) for estado in orden}

    if _quiere_excel():
        return _excel(
            "Conciliación",
            ["Código", "Artículo", "Stock en la app", "Saldo en Sizfra", "Diferencia (app - Sizfra)", "Estado"],
            [list(f.values()) for f in filas],
            "conciliacion-sizfra",
        )
    return jsonify(resumen=resumen, columnas=columnas, filas=filas)
